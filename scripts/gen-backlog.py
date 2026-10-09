"""Regenerate docs/source/backlog.md from docs/source/Nexus_Phase1_Sprint_Plan.xlsx.

Usage (from the repo root):
    python scripts/gen-backlog.py
    python scripts/gen-backlog.py <sheet.xlsx> <preamble.md> <out.md>   # explicit paths

Needs openpyxl (`python -m pip install --user openpyxl`). The sheet is read only.
The hand-written preamble (everything above "## Sprints") is kept from the current
backlog.md; the Sprints table and the Tasks come from the "Sprint plan" and "Backlog"
tabs. Empty fields are left out, and Status is shown only when it is not "Not started".
Check the result with `git diff docs/source/backlog.md`: only the rows you changed
in the sheet should differ.
"""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "source"

if len(sys.argv) == 4:
    xlsx, preamble_src, out = (Path(p) for p in sys.argv[1:4])
elif len(sys.argv) == 1:
    xlsx = SOURCE / "Nexus_Phase1_Sprint_Plan.xlsx"
    preamble_src = out = SOURCE / "backlog.md"
else:
    sys.exit(__doc__)

text = preamble_src.read_text(encoding="utf-8").replace("\r\n", "\n")
preamble = text[: text.index("## Sprints\n")]

wb = openpyxl.load_workbook(xlsx)  # only text columns are read; formulas are not needed


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


lines = [preamble.rstrip("\n"), "", "## Sprints", "",
         "| Sprint | Name | Weeks | Goal | External dependency |",
         "|---|---|---|---|---|"]
for row in wb["Sprint plan"].iter_rows(min_row=2, values_only=True):
    if not row[0] or not str(row[0]).startswith("S"):
        continue  # skip blank and total rows
    lines.append("| " + " | ".join(s(c) for c in row[:5]) + " |")

lines += ["", "## Tasks", ""]
backlog = wb["Backlog"]
col = {c.value: i for i, c in enumerate(backlog[1])}
for row in backlog.iter_rows(min_row=2, values_only=True):
    if not row[0] or "-" not in str(row[0]):
        continue  # skip blank and total rows

    def g(header):
        return s(row[col[header]])

    status = g("Status")
    lines.append(f"### {g('ID')} · {g('Task')}")
    fields = [
        ("Sprint", g("Sprint")),
        ("Week", g("Week")),
        ("Work package", g("Work package")),
        ("Steps (how to do it)", g("Steps (how to do it)")),
        ("Output when done", g("Output when done")),
        ("Now done by", g("Now done by")),
        ("Requirement", g("Requirement")),
        ("Depends on", g("Depends on")),
        ("External blocker", g("External blocker")),
        ("Estimate (days)", g("Estimate (days)")),
        ("Status", "" if status == "Not started" else status),
        ("Acceptance check", g("Acceptance check")),
    ]
    for label, value in fields:
        if value:
            lines.append(f"- **{label}:** " + value.replace("\n", "\n  "))
    lines.append("")

with open(out, "w", encoding="utf-8", newline="\n") as f:
    f.write("\n".join(lines).rstrip("\n") + "\n")
print(f"wrote {out}")
